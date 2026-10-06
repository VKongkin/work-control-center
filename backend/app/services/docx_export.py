"""Turn a knowledge article back into a Word document.

The mirror of `docx_import`, and deliberately its mirror: the importer reads
Word's *semantic* styles - Heading 1, List Bullet, a paragraph style literally
called "Code" - so this writes exactly those names back. That is the whole
trick behind a round trip. A code block exported as "Consolas 9.5pt" would come
home as prose; exported as a paragraph in a style named Code, it comes home as
a code block, because that is the name the importer's style map looks for.

Why generate a document at all, when the imported original is still attached?
Because the original is provenance, not content. The article has been corrected
since - that is the reason for keeping runbooks as text in the first place - so
handing back the vendor's untouched file would hand back the version that was
half wrong. This exports what the article says today.

What survives, in both directions: headings, bullets, numbered steps, bold,
italic, inline code, code blocks, quotes, links, tables with a real header row,
and the images - which are fetched back out of the attachment store and
embedded, so the document someone emails has its screenshots in it rather than
a row of broken references.

One known loss, measured rather than assumed: a second-level bullet is written
as Word's List Bullet 2, which is right on the page and indents properly, but
Word carries list *nesting* in its numbering definitions rather than in the
style name - so importing the document back flattens that sub-bullet to the
top level. Writing the numbering by hand was tried and made it worse: the
paragraph stopped being a list item at all. A one-level flattening on a round
trip is a better bug than a bullet that turns into a sentence.
"""
import io
import re
from typing import Callable, Dict, List, Optional, Tuple

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# A4/Letter with the default margins leaves about six inches of text. An image
# wider than that is not "large", it is off the page.
MAX_IMAGE_WIDTH = Inches(6.0)

# url -> (bytes, content_type). Injected, so this module knows nothing about
# attachments or the database and can be tested on its own.
ImageLoader = Callable[[str], Optional[Tuple[bytes, str]]]

FENCE = re.compile(r"^\s*(```|~~~)\s*([A-Za-z0-9_+-]*)\s*$")
HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
RULE = re.compile(r"^\s*([-*_])(?:\s*\1){2,}\s*$")
BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
NUMBER = re.compile(r"^(\s*)\d+[.)]\s+(.*)$")
QUOTE = re.compile(r"^\s*>\s?(.*)$")
TABLE_ROW = re.compile(r"^\s*\|(.+)\|\s*$")
TABLE_RULE = re.compile(r"^\s*\|[\s:|-]+\|\s*$")

# One pass over a line of text. Images and links first so their contents are
# not mistaken for emphasis, then code, then bold before italic - `**` has to
# win over `*` or every bold run becomes two italic ones around a star.
INLINE = re.compile(
    r"!\[(?P<alt>[^\]]*)\]\((?P<src>[^)\s]+)(?:\s+\"[^\"]*\")?\)"
    r"|\[(?P<text>[^\]]+)\]\((?P<href>[^)\s]+)(?:\s+\"[^\"]*\")?\)"
    r"|`(?P<code>[^`]+)`"
    r"|\*\*(?P<bold>.+?)\*\*"
    r"|__(?P<bold2>.+?)__"
    r"|(?<!\w)\*(?P<it>[^*]+)\*(?!\w)"
    r"|(?<!\w)_(?P<it2>[^_]+)_(?!\w)"
)


# ------------------------------------------------------------------ styles

def _ensure_styles(doc: Document) -> None:
    """Add the two styles Word's default template does not carry.

    Named, not merely formatted. `docx_import.STYLE_MAP` matches on the name:
    `p[style-name='Code'] => pre` and `r[style-name='Code Char'] => code`. The
    font here is only so it looks right to a human; the name is what makes the
    content survive coming back.
    """
    names = {s.name for s in doc.styles}
    if "Code" not in names:
        st = doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["No Spacing"]
        st.font.name = "Consolas"
        st.font.size = Pt(9.5)
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.left_indent = Inches(0.25)
    if "Code Char" not in names:
        ch = doc.styles.add_style("Code Char", WD_STYLE_TYPE.CHARACTER)
        ch.font.name = "Consolas"
        ch.font.size = Pt(9.5)


def _hyperlink(paragraph, url: str, text: str) -> None:
    """A real Word hyperlink.

    python-docx has no API for this, so the relationship and the w:hyperlink
    element are written by hand. Worth the twelve lines: a runbook is mostly
    links to tickets and consoles, and a link exported as grey text is a link
    nobody can follow.
    """
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = paragraph._p.makeelement(qn("w:hyperlink"), {qn("r:id"): r_id})
    run = paragraph.add_run(text)
    run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
    run.font.underline = True
    link.append(run._r)
    paragraph._p.append(link)


def _picture(paragraph, data: bytes, load_failed_alt: str) -> bool:
    try:
        run = paragraph.add_run()
        pic = run.add_picture(io.BytesIO(data))
    except Exception:
        # A format Word will not take (an EMF kept as an attachment, say).
        # The alt text is better than a hole in the page.
        paragraph.add_run(f"[{load_failed_alt}]").italic = True
        return False
    if pic.width > MAX_IMAGE_WIDTH:
        pic.height = int(pic.height * (MAX_IMAGE_WIDTH / pic.width))
        pic.width = MAX_IMAGE_WIDTH
    return True


# ------------------------------------------------------------------ inline

def _emit(paragraph, text: str, load_image: Optional[ImageLoader],
          bold=False, italic=False) -> None:
    """Write one line of markdown into a paragraph, as styled runs."""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            _plain(paragraph, text[pos:m.start()], bold, italic)
        g = m.groupdict()
        if g["src"] is not None:
            loaded = load_image(g["src"]) if load_image else None
            if loaded:
                _picture(paragraph, loaded[0], g["alt"] or "image")
            else:
                _plain(paragraph, f"[{g['alt'] or 'image'}]", bold, True)
        elif g["href"] is not None:
            _hyperlink(paragraph, g["href"], g["text"])
        elif g["code"] is not None:
            run = paragraph.add_run(g["code"])
            run.style = paragraph.part.document.styles["Code Char"]
            run.bold, run.italic = bold, italic
        elif g["bold"] is not None or g["bold2"] is not None:
            _emit(paragraph, g["bold"] or g["bold2"], load_image, True, italic)
        else:
            _emit(paragraph, g["it"] or g["it2"], load_image, bold, True)
        pos = m.end()
    if pos < len(text):
        _plain(paragraph, text[pos:], bold, italic)


def _plain(paragraph, text: str, bold: bool, italic: bool) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    run.bold, run.italic = bold, italic


# ------------------------------------------------------------------ blocks

def _table(doc: Document, rows: List[List[str]], load_image) -> None:
    """A table whose first row is a header Word actually knows about.

    Two deliberate choices, both learned by watching what comes back:

    The row is marked `w:tblHeader`, which makes Word repeat it across a page
    break and makes the importer see a `<th>` - so the markdown gets a real
    header row rather than the empty one `promote_empty_table_header` exists to
    rescue.

    The header is *not* bolded run by run. A table style whose own conditional
    formatting bolds the first row looks identical in Word, but carries no bold
    property on the text itself - so the header comes home as `| Node |` rather
    than as `| **Node** |`, with the emphasis markers Word put there.
    """
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = doc.styles["Light Grid"]
    for y, row in enumerate(rows):
        for x in range(width):
            cell = table.cell(y, x)
            cell.paragraphs[0].text = ""
            _emit(cell.paragraphs[0], row[x] if x < len(row) else "", load_image)
    tr = table.rows[0]._tr
    tr.get_or_add_trPr().append(tr.makeelement(qn("w:tblHeader"), {}))


def _split_row(line: str) -> List[str]:
    inner = TABLE_ROW.match(line).group(1)
    return [c.strip() for c in inner.split("|")]


def _rule(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _plain(p, "* * *", False, False)


def render(doc: Document, markdown: str, load_image: Optional[ImageLoader] = None) -> None:
    """Write a markdown body into an open document."""
    lines = (markdown or "").replace("\r\n", "\n").split("\n")
    i = 0
    para: List[str] = []

    def flush():
        nonlocal para
        if para:
            _emit(doc.add_paragraph(), " ".join(para), load_image)
            para = []

    while i < len(lines):
        line = lines[i]

        fence = FENCE.match(line)
        if fence:
            flush()
            i += 1
            body = []
            while i < len(lines) and not FENCE.match(lines[i]):
                body.append(lines[i])
                i += 1
            i += 1                                  # the closing fence
            for text in body or [""]:
                # One paragraph per line: Word has no multi-line paragraph that
                # keeps its line breaks through an import, and a shell session
                # collapsed onto one line is not a shell session.
                doc.add_paragraph(text, style="Code")
            continue

        if not line.strip():
            flush()
            i += 1
            continue

        m = HEADING.match(line)
        if m:
            flush()
            _emit(doc.add_heading("", level=min(6, len(m.group(1)))), m.group(2), load_image)
            i += 1
            continue

        if RULE.match(line):
            flush()
            _rule(doc)
            i += 1
            continue

        if TABLE_ROW.match(line) and i + 1 < len(lines) and TABLE_RULE.match(lines[i + 1]):
            flush()
            rows = [_split_row(line)]
            i += 2
            while i < len(lines) and TABLE_ROW.match(lines[i]):
                rows.append(_split_row(lines[i]))
                i += 1
            _table(doc, rows, load_image)
            continue

        m = QUOTE.match(line)
        if m:
            flush()
            said = [m.group(1)]
            i += 1
            while i < len(lines) and QUOTE.match(lines[i]):
                said.append(QUOTE.match(lines[i]).group(1))
                i += 1
            _emit(doc.add_paragraph(style="Quote"), " ".join(s for s in said if s), load_image)
            continue

        m = BULLET.match(line) or NUMBER.match(line)
        if m:
            flush()
            ordered = NUMBER.match(line) is not None
            depth = 2 if len(m.group(1)) >= 2 else 1
            style = ("List Number" if ordered else "List Bullet") + ("" if depth == 1 else " 2")
            _emit(doc.add_paragraph(style=style), m.group(2), load_image)
            i += 1
            continue

        para.append(line.strip())
        i += 1

    flush()


# ------------------------------------------------------------------- build

def build(title: str, markdown: str, summary: Optional[str] = None,
          properties: Optional[Dict[str, str]] = None,
          load_image: Optional[ImageLoader] = None) -> bytes:
    """One article, as the bytes of a .docx."""
    doc = Document()
    _ensure_styles(doc)

    # Heading 1, not Word's "Title" style, and the reason is the round trip.
    # Mammoth reads Heading 1 and nothing else as the document's own name: with
    # the Title style it comes back as an ordinary paragraph, the importer
    # titles the new article after its first *section* instead, and the old
    # title stays in the body - so every export-import cycle prepends another
    # line of title. Heading 1 is recognised, becomes the article's title, and
    # is then stripped from the body it opens.
    _emit(doc.add_heading("", level=1), title or "Untitled", load_image)
    if summary:
        # Deliberately italic runs rather than Word's Subtitle style, which
        # mammoth does not recognise: an unrecognised style is a warning shown
        # to whoever imports this, and "Unrecognised paragraph style: Subtitle"
        # on a file WCC wrote itself is noise that teaches people to ignore
        # warnings that matter.
        _emit(doc.add_paragraph(), f"*{summary}*", load_image)

    render(doc, markdown, load_image)

    # Kind, status, tags and the rest go in Word's own document properties
    # rather than printed across the top of the page: the person emailing this
    # wants the runbook, not a header block about where it was filed. Word
    # shows them under File > Info, and they survive being mailed about.
    core = doc.core_properties
    core.title = title or "Untitled"
    for key, value in (properties or {}).items():
        if not value:
            continue
        if key == "summary":
            core.subject = value[:255]
        elif key == "tags":
            core.keywords = value[:255]
        elif key == "comments":
            core.comments = value[:4000]
        elif key == "category":
            core.category = value[:64]

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def filename_for(title: str, article_id: int) -> str:
    """A filename that survives every operating system someone might save it on."""
    stem = re.sub(r"[^A-Za-z0-9 _-]+", "", title or "").strip()
    stem = re.sub(r"\s+", " ", stem)[:80].strip()
    return f"{stem or f'article-{article_id}'}.docx"
